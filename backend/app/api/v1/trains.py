from fastapi import APIRouter, Depends, Query, UploadFile, File, HTTPException
from fastapi.responses import StreamingResponse
from typing import Optional, List
from datetime import date, datetime
from sqlalchemy.orm import Session
import csv
import io
from app.core.database import get_db
from app.schemas.train import TrainCreate, TrainOut, TrainImportResult
from app.models.train import Train
from app.services.train_service import TrainService, IMPORT_COLUMNS_DOC
from app.services.ir_integration import mock_schedule_for_section

router = APIRouter()
service = TrainService()


@router.get("", response_model=List[TrainOut])
def list_trains(section: Optional[str] = Query(None), skip: int = 0, limit: int = 200, db: Session = Depends(get_db)):
    return service.list(db, section=section, skip=skip, limit=limit)


@router.post("", response_model=TrainOut, status_code=201)
def create_train(payload: TrainCreate, db: Session = Depends(get_db)):
    obj = Train(**payload.model_dump())
    db.add(obj)
    db.commit()
    db.refresh(obj)
    return obj


@router.post("/seed")
def seed(section: str = "Bhadrak–Jajpur", db: Session = Depends(get_db)):
    n = service.seed_if_empty(db, section=section)
    return {"seeded": n}


@router.get("/live")
async def live_trains(section: str = "Bhadrak–Jajpur", day: Optional[date] = None, db: Session = Depends(get_db)):
    d = day or date.today()
    trains = await service.get_for_section(db, section, datetime(d.year, d.month, d.day))
    # serialise datetimes
    out = []
    for t in trains:
        st = t.get("scheduled_time")
        out.append({**t, "scheduled_time": st.isoformat() if hasattr(st, "isoformat") else st})
    return {"section": section, "date": d.isoformat(), "count": len(out), "trains": out}


def _parse_csv(content: bytes) -> List[dict]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise HTTPException(status_code=400, detail="CSV must be UTF-8 encoded")
    try:
        reader = csv.DictReader(io.StringIO(text))
        if not reader.fieldnames:
            raise HTTPException(status_code=400, detail="CSV has no header row")
        return [dict(r) for r in reader if any((v or "").strip() for v in r.values())]
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Could not parse CSV: {e}")


def _parse_xlsx(content: bytes) -> List[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise HTTPException(status_code=400, detail="XLSX support not installed (openpyxl missing)")
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Could not parse XLSX: {e}")
    if not rows or not any(rows[0]):
        raise HTTPException(status_code=400, detail="XLSX sheet is empty")
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    out = []
    for r in rows[1:]:
        vals = [("" if v is None else v) for v in r]
        if not any(str(v).strip() for v in vals):
            continue
        out.append(dict(zip(headers, vals)))
    return out


@router.get("/columns")
def import_columns():
    """Document the expected timetable file format for the upload UI."""
    return {
        "formats": ["csv", "xlsx"],
        "columns": IMPORT_COLUMNS_DOC,
        "required": ["train_number", "train_name", "scheduled_time"],
        "direction_note": "Direction is derived as origin → destination.",
    }


@router.get("/sample")
def sample_csv(day: Optional[date] = None, section: str = "Bhadrak–Jajpur"):
    """Generate a ready-to-upload sample timetable for the given date (default today)."""
    d = day or date.today()
    rows = mock_schedule_for_section(section, d)
    buf = io.StringIO()
    writer = csv.DictWriter(
        buf,
        fieldnames=["train_number", "train_name", "train_type", "priority", "section",
                    "origin", "destination", "scheduled_time", "status"],
    )
    writer.writeheader()
    for m in rows:
        writer.writerow({
            "train_number": m["train_number"],
            "train_name": m["train_name"],
            "train_type": m["train_type"],
            "priority": m["priority"],
            "section": m["section"],
            "origin": m["origin"],
            "destination": m["destination"],
            "scheduled_time": m["scheduled_time"].isoformat(timespec="minutes"),
            "status": m["status"],
        })
    buf.seek(0)
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename=rail-sway-sample-timetable-{d.isoformat()}.csv"},
    )


@router.post("/import", response_model=TrainImportResult)
async def import_timetable(file: UploadFile = File(...), db: Session = Depends(get_db)):
    name = (file.filename or "").lower()
    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="Uploaded file is empty")
    if name.endswith(".csv"):
        rows = _parse_csv(content)
    elif name.endswith((".xlsx", ".xlsm")):
        rows = _parse_xlsx(content)
    else:
        raise HTTPException(status_code=400, detail="Only .csv and .xlsx files are supported")
    if not rows:
        raise HTTPException(status_code=400, detail="No data rows found in file")
    return service.import_timetable(db, rows)
